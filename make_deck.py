from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Initialize presentation
prs = Presentation()

# Define slide layout (1 = Title and Content)
slide_layout = prs.slide_layouts[1] 

# Slide 1: Title Slide
title_slide = prs.slides.add_slide(prs.slide_layouts[0])
title_slide.shapes.title.text = "Your Presentation Title"
title_slide.placeholders[1].text = "Subtitle or Author Name"

# Slide 2: Bullet Points
slide_2 = prs.slides.add_slide(slide_layout)
slide_2.shapes.title.text = "Key Objectives"
tf_2 = slide_2.placeholders[1].text_frame
tf_2.text = "First main takeaway or point"
tf_2.add_paragraph().text = "Second main takeaway or point"
tf_2.add_paragraph().text = "Third main takeaway or point"

# Save presentation
prs.save("presentation.pptx")
print("Saved as presentation.pptx!")